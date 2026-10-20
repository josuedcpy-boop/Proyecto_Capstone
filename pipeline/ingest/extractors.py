"""Extractores de texto plano por tipo de archivo. Cada uno devuelve "" si no hay texto útil."""

from __future__ import annotations

from pathlib import Path


class UnsupportedFileError(Exception):
    pass


# Extrae el texto de un .docx: párrafos y, aparte, el contenido de cualquier tabla.
def extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# Extrae el texto de un .pptx: recorre cada diapositiva y cada cuadro de texto.
def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts)


# Extrae el texto de un .xlsx: cada fila de cada hoja se junta en una línea "celda | celda".
def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# Extrae el texto de un .pdf (solo funciona si el PDF tiene capa de texto real, no un
# escaneo de imagen puro — en ese caso cada página devuelve "").
def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(p for p in parts if p.strip())


# Lee un .txt tal cual (ej. los capítulos de libro ya convertidos de PDF a texto).
def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


EXTRACTORS = {
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".pdf": extract_pdf,
    ".txt": extract_txt,
}


def extract_text(path: Path) -> str:
    """Despacha al extractor según la extensión. Lanza UnsupportedFileError si no hay soporte."""
    extractor = EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise UnsupportedFileError(f"Sin extractor para {path.suffix}: {path.name}")
    return extractor(path)
