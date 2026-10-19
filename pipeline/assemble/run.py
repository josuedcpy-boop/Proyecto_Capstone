"""Ensambla las preguntas validadas en un examen final .docx (abrible en Word y similares)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import docx
from docx.shared import Pt

from pipeline.assemble.naming import build_output_path, clean_filename_part
from pipeline.config import EXAMEN_GENERADO_DIR
from pipeline.schemas import GeneratedExam

_SUBJECT_LABELS = {"history": "Historia", "technology": "Tecnología"}


def _add_heading(document: docx.Document, text: str, size: int, bold: bool = True) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def assemble_exam(exam: GeneratedExam, tema: str) -> Path:
    exam_path = build_output_path(EXAMEN_GENERADO_DIR, tema, exam.grade)
    tema_clean = clean_filename_part(tema) or "Examen"

    subject_label = _SUBJECT_LABELS.get(exam.subject, exam.subject)
    document = docx.Document()

    _add_heading(document, f"Examen generado — {tema_clean} ({subject_label}, grado {exam.grade})", size=16)
    document.add_paragraph(f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")

    errors = [i for i in exam.issues if i.severity == "error"]
    warnings = [i for i in exam.issues if i.severity == "warning"]
    if errors:
        warning_paragraph = document.add_paragraph()
        warning_run = warning_paragraph.add_run(
            f"⚠ {len(errors)} problema(s) crítico(s) detectados en validación — revisar antes de usar en clase "
            "(ver detalle al final del documento)."
        )
        warning_run.bold = True

    document.add_paragraph()

    for n, question in enumerate(exam.questions, start=1):
        flagged = any(i.question_id == question.id for i in exam.issues)
        marker = " ⚠" if flagged else ""
        meta_paragraph = document.add_paragraph()
        meta_run = meta_paragraph.add_run(
            f"{n}. ({question.bloom_level or '?'} / DOK{question.dok_level or '?'}){marker}"
        )
        meta_run.italic = True
        document.add_paragraph(question.text)
        document.add_paragraph()

    if exam.issues:
        document.add_page_break()
        _add_heading(document, "Reporte de validación", size=13)
        for issue in errors + warnings:
            document.add_paragraph(
                f"[{issue.severity}] {issue.question_id} ({issue.rule}): {issue.detail}",
                style="List Bullet",
            )

    document.save(str(exam_path))
    return exam_path
