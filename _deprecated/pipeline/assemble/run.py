"""Ensambla las preguntas validadas en un examen final .docx (abrible en Word y similares)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import docx
from docx.shared import Pt

from pipeline.assemble.naming import build_output_path, clean_filename_part
from pipeline.config import EXAMEN_GENERADO_DIR
from pipeline.schemas import AuditReport, GeneratedExam

_OPTION_LETTERS = "abcd"

_SUBJECT_LABELS = {"history": "Historia", "technology": "Tecnología"}


# Agrega un párrafo en negrita a modo de título dentro del documento.
def _add_heading(document: docx.Document, text: str, size: int, bold: bool = True) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


# Construye el .docx final: título, cada pregunta con su tipo/puntos (y opciones si es de
# opción múltiple), y al final el reporte de validación y de auditoría si se dieron.
def assemble_exam(
    exam: GeneratedExam, tema: str, audit: AuditReport | None = None, reference_path: str | None = None
) -> Path:
    exam_path = build_output_path(EXAMEN_GENERADO_DIR, tema, exam.grade)
    tema_clean = clean_filename_part(tema) or "Examen"

    subject_label = _SUBJECT_LABELS.get(exam.subject, exam.subject)
    document = docx.Document()

    _add_heading(document, f"Examen generado — {tema_clean} ({subject_label}, grado {exam.grade})", size=16)
    document.add_paragraph(f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    if reference_path:
        # Guardado para que `pipeline.cli audit` pueda re-auditar este examen después sin
        # que el usuario tenga que acordarse de cuál fue el material de referencia.
        document.add_paragraph(f"Material de referencia: {reference_path}")

    errors = [i for i in exam.issues if i.severity == "error"]
    warnings = [i for i in exam.issues if i.severity == "warning"]
    if errors:
        warning_paragraph = document.add_paragraph()
        warning_run = warning_paragraph.add_run(
            f"⚠ {len(errors)} problema(s) crítico(s) detectados en validación — revisar antes de usar en clase "
            "(ver detalle al final del documento)."
        )
        warning_run.bold = True

    document.add_paragraph()

    for n, question in enumerate(exam.questions, start=1):
        flagged = any(i.question_id == question.id for i in exam.issues)
        marker = " ⚠" if flagged else ""
        meta_paragraph = document.add_paragraph()
        meta_run = meta_paragraph.add_run(
            f"{n}. ({question.bloom_level or '?'} / DOK{question.dok_level or '?'} / "
            f"{question.part_type or '?'} / {question.points or 0}pts){marker}"
        )
        meta_run.italic = True
        document.add_paragraph(question.text)
        if question.options:
            for letter, option in zip(_OPTION_LETTERS, question.options):
                document.add_paragraph(f"   {letter}) {option}")
        document.add_paragraph()

    if exam.issues:
        document.add_page_break()
        _add_heading(document, "Reporte de validación", size=13)
        for issue in errors + warnings:
            document.add_paragraph(
                f"[{issue.severity}] {issue.question_id} ({issue.rule}): {issue.detail}",
                style="List Bullet",
            )

    if audit is not None:
        document.add_page_break()
        estado = "APROBADO" if audit.aprobado_general else "NO APROBADO"
        _add_heading(document, f"Reporte de auditoría — {estado}", size=13)
        for criterio in audit.criterios:
            marca = "[OK]" if criterio.aprobado else "[FALLA]"
            document.add_paragraph(f"{marca} {criterio.criterio}: {criterio.detalle}", style="List Bullet")
            if criterio.explicacion:
                document.add_paragraph(f"Por qué falla: {criterio.explicacion}", style="List Bullet 2")

    document.save(str(exam_path))
    return exam_path
