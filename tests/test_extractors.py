from pathlib import Path

import docx

from pipeline.ingest.extractors import UnsupportedFileError, extract_text


def test_extract_docx_roundtrip(tmp_path: Path):
    doc = docx.Document()
    doc.add_paragraph("1. ¿Cuál es la capital de Puerto Rico?")
    doc.add_paragraph("2. Explica el Grito de Lares.")
    path = tmp_path / "examen.docx"
    doc.save(str(path))

    text = extract_text(path)
    assert "capital de Puerto Rico" in text
    assert "Grito de Lares" in text


def test_unsupported_extension_raises(tmp_path: Path):
    path = tmp_path / "archivo.xyz"
    path.write_text("contenido")
    try:
        extract_text(path)
        assert False, "debería haber lanzado UnsupportedFileError"
    except UnsupportedFileError:
        pass
